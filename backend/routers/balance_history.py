import io
import logging

import openpyxl
import requests
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from core.deps import get_current_user
from db import cashout as cashout_db
from db import balance_history as bh_db
from db import users as users_db
from services.encryption import decrypt_value

log = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/bh")

EXCEL_HEADERS = [
    "Счёт", "accountId", "Дата",
    "totalAssets", "cashBalance", "portfolioValue", "blocked", "blockedForOrders",
    "assetMarketValue", "cryptoBalance", "futureCashFlow", "unrealizedPnl",
    "accruedInterest", "totalUnrealizedPnl",
    "marginUtilization", "marginUsage", "marginBalance", "marginAvailable",
    "positionMargin", "cashMargin", "orderMargin",
    "cashIn", "cashOut", "cashAvailable", "externalTransfers",
    "dailyPnl", "dailyPnlPercent", "prevDayTotalAssets",
]

_FIELD_KEYS = [
    None, None, None,
    "totalAssets", "cashBalance", "portfolioValue",
    "blocked", "blockedForOrders", "assetMarketValue",
    "cryptoBalance", "futureCashFlow", "unrealizedPnl",
    "accruedInterest", "totalUnrealizedPnl",
    "marginUtilization", "marginUsage", "marginBalance", "marginAvailable",
    "positionMargin", "cashMargin", "orderMargin",
    "cashIn", "cashOut", "cashAvailable", "externalTransfers",
    "dailyPnl", "dailyPnlPercent", "prevDayTotalAssets",
]

_ROW_FILL = PatternFill("solid", fgColor="D6E4F0")  # светло-синий для всех строк

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_THIN = Side(style="thin", color="B0BEC5")
_THICK = Side(style="medium", color="455A64")


def _border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN) -> Border:
    return Border(left=left, right=right, top=top, bottom=bottom)


def _apply_header(ws) -> None:
    for col, header in enumerate(EXCEL_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
    ws.row_dimensions[1].height = 40


def _auto_widths(ws, min_w: int = 10, max_w: int = 30) -> None:
    for col_cells in ws.columns:
        col_idx = col_cells[0].column
        max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(min_w, min(max_len + 2, max_w))


class BHAccountCreate(BaseModel):
    name: str
    account_id: str
    asset_id: str


def _unity_cfg(current_user: str) -> tuple[str, str]:
    user = users_db.get_user_by_username(current_user)
    if not user:
        raise HTTPException(401, "Пользователь не найден")
    cfg = cashout_db.get_unity_config(user.id)
    base_url = (cfg.get("base_url") or "").rstrip("/")
    if not base_url:
        raise HTTPException(400, "Unity API не настроен. Укажите Base URL в настройках.")
    token = decrypt_value(cfg["auth_token_enc"]) if cfg.get("auth_token_enc") else ""
    return base_url, token


def _fetch(base_url: str, token: str, account_id: str, asset_id: str, from_date: str, to_date: str) -> list:
    url = f"{base_url}/api/v1/balanceHistory"
    headers = {"auth-token": token} if token else {}
    params = {"accountId": account_id, "assetId": asset_id, "from": from_date, "to": to_date}
    try:
        resp = requests.get(url, headers=headers, params=params, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        return data if isinstance(data, list) else []
    except Exception as e:
        log.warning("Unity balanceHistory error account=%s: %s", account_id, e)
        return []


# ── Счета ─────────────────────────────────────────────────────────────────────

@router.get("/accounts")
async def list_accounts(current_user: str = Depends(get_current_user)):
    return bh_db.list_bh_accounts()


@router.post("/accounts", status_code=201)
async def create_account(req: BHAccountCreate, current_user: str = Depends(get_current_user)):
    result = bh_db.create_bh_account(req.name, req.account_id, req.asset_id)
    if not result:
        raise HTTPException(500, "Ошибка создания счёта")
    return result


@router.delete("/accounts/{row_id}")
async def delete_account(row_id: int, current_user: str = Depends(get_current_user)):
    if not bh_db.delete_bh_account(row_id):
        raise HTTPException(500, "Ошибка удаления")
    return {"ok": True}


# ── Данные ────────────────────────────────────────────────────────────────────

@router.get("/data")
async def get_data(
    from_date: str = Query(..., alias="from"),
    to_date: str = Query(..., alias="to"),
    current_user: str = Depends(get_current_user),
):
    base_url, token = _unity_cfg(current_user)
    rows = []
    for acc in bh_db.list_bh_accounts():
        for entry in _fetch(base_url, token, acc["account_id"], acc["asset_id"], from_date, to_date):
            rows.append({"account_name": acc["name"], "account_id": acc["account_id"], **entry})
    return rows


@router.get("/export")
async def export_excel(
    from_date: str = Query(..., alias="from"),
    to_date: str = Query(..., alias="to"),
    current_user: str = Depends(get_current_user),
):
    base_url, token = _unity_cfg(current_user)
    accounts = bh_db.list_bh_accounts()

    wb = openpyxl.Workbook()

    # ── Общий лист со всеми счетами ──────────────────────────────────────────
    ws_all = wb.active
    ws_all.title = "Все счета"
    _apply_header(ws_all)
    ws_all.freeze_panes = "D2"  # фиксируем шапку + первые 3 колонки
    ws_all.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_HEADERS))}1"

    num_fmt = '#,##0.00'
    date_fmt = 'DD.MM.YYYY'

    all_row = 2
    prev_account_id = None

    for acc in accounts:
        entries = _fetch(base_url, token, acc["account_id"], acc["asset_id"], from_date, to_date)
        if not entries:
            continue

        is_new_group = prev_account_id != acc["account_id"]

        for entry in entries:
            b = entry.get("balance", {})
            row_values = [
                acc["name"], acc["account_id"], entry.get("date"),
                b.get("totalAssets"), b.get("cashBalance"), b.get("portfolioValue"),
                b.get("blocked"), b.get("blockedForOrders"), b.get("assetMarketValue"),
                b.get("cryptoBalance"), b.get("futureCashFlow"), b.get("unrealizedPnl"),
                b.get("accruedInterest"), b.get("totalUnrealizedPnl"),
                b.get("marginUtilization"), b.get("marginUsage"), b.get("marginBalance"),
                b.get("marginAvailable"), b.get("positionMargin"), b.get("cashMargin"),
                b.get("orderMargin"), b.get("cashIn"), b.get("cashOut"),
                b.get("cashAvailable"), b.get("externalTransfers"),
                b.get("dailyPnl"), b.get("dailyPnlPercent"), b.get("prevDayTotalAssets"),
            ]

            top_side = _THICK if is_new_group else _THIN
            for col_idx, val in enumerate(row_values, 1):
                cell = ws_all.cell(row=all_row, column=col_idx, value=val)
                cell.fill = _ROW_FILL
                cell.border = _border(top=top_side)
                if col_idx == 3 and val:  # Дата
                    cell.number_format = date_fmt
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx > 3 and isinstance(val, (int, float)):  # числа
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx <= 2:
                    cell.alignment = Alignment(horizontal="left")
            is_new_group = False
            prev_account_id = acc["account_id"]
            all_row += 1

    _auto_widths(ws_all)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"balances_{from_date}_{to_date}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
