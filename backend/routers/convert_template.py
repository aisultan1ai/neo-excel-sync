import io
import logging
import re
import uuid
import zipfile
from datetime import datetime
from threading import Lock
from typing import Any, Dict

import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.deps import get_current_user
from utils.files import cleanup_files, save_upload_file

log = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/convert-template")

_ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}
_MAX_UPLOAD_BYTES = 50 * 1024 * 1024

_CACHE: Dict[str, Any] = {}
_CACHE_LOCK = Lock()

REQUIRED_COLUMNS = [
    "Trade ID", "Symbol", "Side", "Price", "Quantity", "Amount", "Fee", "Time",
]
TEMPLATE_COLUMNS = [
    "ID", "Instrument", "Amount", "Quote amount", "Price", "Side",
    "Commission", "Trade type", "Trade date", "Value date",
    "Transact time", "Account", "Closed P/L", "Net commission amount", "Absolute amount",
]
SHEET_NAME = "отчет по сделкам (3)"

_DATE_RE     = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")   # DD.MM.YYYY
_DATE_ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")    # YYYY-MM-DD (pandas auto-parse)


def _normalize_date(date_part: str) -> str:
    """Convert YYYY-MM-DD → DD.MM.YYYY. Returns input unchanged for other formats."""
    if _DATE_ISO_RE.match(date_part):
        y, m, d = date_part.split("-")
        return f"{d}.{m}.{y}"
    return date_part


# ── Validation ─────────────────────────────────────────────────

def _validate_raw_row(row: pd.Series, seen_ids: set) -> list:
    issues = []

    # Duplicate Trade ID
    trade_id = str(row.get("Trade ID", "")).strip()
    if trade_id in seen_ids:
        issues.append({"type": "error", "field": "Trade ID",
                        "msg": f"Дублирующийся Trade ID: «{trade_id}»"})
    seen_ids.add(trade_id)

    # Side must be BUY or SELL
    side = str(row.get("Side", "")).strip().upper()
    if side not in ("BUY", "SELL"):
        issues.append({"type": "error", "field": "Side",
                        "msg": f"Некорректный Side: «{side}» (ожидается BUY или SELL)"})

    # Numeric integrity: Price, Quantity, Amount
    for field in ("Price", "Quantity", "Amount"):
        raw = row.get(field)
        try:
            val = float(raw)
            if val < 0:
                issues.append({"type": "error", "field": field,
                                "msg": f"{field} отрицательный: {val}"})
        except (ValueError, TypeError):
            issues.append({"type": "error", "field": field,
                            "msg": f"{field} не является числом: «{raw}»"})

    # Fee: must be numeric (negative values are allowed)
    fee_raw = str(row.get("Fee", "")).strip()
    fee_clean = re.sub(r"\s+[A-Z]+$", "", fee_raw).strip()
    try:
        float(fee_clean)
    except (ValueError, TypeError):
        issues.append({"type": "error", "field": "Fee",
                        "msg": f"Fee не является числом: «{fee_raw}»"})

    # Date format in Time column
    time_str = str(row.get("Time", "")).strip()
    parts = time_str.split()
    date_part = _normalize_date(parts[0] if parts else "")
    if not _DATE_RE.match(date_part):
        issues.append({"type": "warning", "field": "Time",
                        "msg": f"Нестандартный формат даты: «{date_part}» (ожидается ДД.ММ.ГГГГ)"})

    # Price × Quantity ≈ Amount (tolerance 0.01%)
    try:
        price = float(row.get("Price", 0))
        qty = float(row.get("Quantity", 0))
        amount = float(row.get("Amount", 0))
        if amount != 0:
            expected = price * qty
            deviation = abs(expected - abs(amount)) / abs(amount)
            if deviation > 0.0001:
                issues.append({"type": "warning", "field": "Amount",
                                "msg": (f"Price×Qty={expected:.6f} ≠ Amount={amount:.6f} "
                                        f"(откл. {deviation * 100:.4f}%)")})
    except (ValueError, TypeError):
        pass

    return issues


# ── Transformation ─────────────────────────────────────────────

def _transform_symbol(symbol: str, exchange: str) -> str:
    return f"FU.{symbol.strip()}.{exchange.strip().upper()}.Z2099"


def _transform_fee(fee) -> str:
    cleaned = re.sub(r"\s+[A-Z]+$", "", str(fee).strip()).strip()
    try:
        return f"{float(cleaned):.10f}".rstrip("0").rstrip(".")
    except (ValueError, TypeError):
        return cleaned


def _split_datetime(dt_str: str):
    parts = str(dt_str).strip().split()
    date_part = _normalize_date(parts[0] if parts else "")
    time_part = parts[1] if len(parts) > 1 else ""
    return date_part, time_part


def _transform_row(row: pd.Series, account: str, exchange: str) -> dict:
    date_part, time_part = _split_datetime(row["Time"])
    transact_time = f"{date_part}  {time_part}".strip() if time_part else date_part
    return {
        "ID": str(row["Trade ID"]),
        "Instrument": _transform_symbol(str(row["Symbol"]), exchange),
        "Amount": row["Quantity"],
        "Quote amount": row["Amount"],
        "Price": row["Price"],
        "Side": str(row["Side"]).strip().upper(),
        "Commission": _transform_fee(row["Fee"]),
        "Trade type": "DIRECT",
        "Trade date": date_part,
        "Value date": "31.12.2099",
        "Transact time": transact_time,
        "Account": account,
        "Closed P/L": "",
        "Net commission amount": "USDT",
        "Absolute amount": row["Quantity"],
    }


def _load_file(filepath: str, ext: str) -> pd.DataFrame:
    if ext == ".csv":
        try:
            df = pd.read_csv(filepath, sep=";", encoding="utf-8-sig")
            if len(df.columns) < 2:
                df = pd.read_csv(filepath, sep=",", encoding="utf-8-sig")
        except Exception as e:
            raise ValueError(f"Не удалось разобрать CSV: {e}")
    else:
        try:
            df = pd.read_excel(filepath)
        except Exception as e:
            raise ValueError(f"Не удалось разобрать Excel: {e}")
    return df


def _process_file(filepath: str, ext: str, account: str, exchange: str) -> dict:
    df = _load_file(filepath, ext)
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Отсутствуют колонки: {', '.join(missing)}")

    # Validate all rows before transformation
    seen_ids: set = set()
    all_issues = []
    for _, row in df.iterrows():
        all_issues.append(_validate_raw_row(row, seen_ids))

    # Transform
    rows = [_transform_row(row, account, exchange) for _, row in df.iterrows()]
    result_df = pd.DataFrame(rows, columns=TEMPLATE_COLUMNS)

    instruments = int(result_df["Instrument"].nunique())
    dates = result_df["Trade date"].replace("", pd.NA).dropna().unique().tolist()
    if dates:
        try:
            sorted_dates = sorted(dates, key=lambda d: datetime.strptime(d, "%d.%m.%Y"))
            date_range = f"{sorted_dates[0]} — {sorted_dates[-1]}"
        except Exception:
            date_range = f"{min(dates)} — {max(dates)}"
    else:
        date_range = "—"

    # Build preview with severity markers
    preview_records = result_df.head(10).fillna("").to_dict(orient="records")
    for i, row_dict in enumerate(preview_records):
        issues = all_issues[i] if i < len(all_issues) else []
        has_error = any(iss["type"] == "error" for iss in issues)
        has_warn = any(iss["type"] == "warning" for iss in issues)
        row_dict["_issues"] = issues
        row_dict["_severity"] = "error" if has_error else ("warning" if has_warn else None)

    # Build full validation report
    rows_detail = []
    for i, issues in enumerate(all_issues):
        if not issues:
            continue
        trade_id = str(df.iloc[i]["Trade ID"]) if "Trade ID" in df.columns else str(i + 1)
        has_error = any(iss["type"] == "error" for iss in issues)
        rows_detail.append({
            "row_num": i + 1,
            "trade_id": trade_id,
            "severity": "error" if has_error else "warning",
            "issues": issues,
        })

    total_errors = sum(
        1 for issues in all_issues for iss in issues if iss["type"] == "error"
    )
    total_warnings = sum(
        1 for issues in all_issues for iss in issues if iss["type"] == "warning"
    )
    rows_with_errors = sum(
        1 for issues in all_issues if any(iss["type"] == "error" for iss in issues)
    )
    rows_with_warnings = sum(
        1 for issues in all_issues
        if any(iss["type"] == "warning" for iss in issues)
        and not any(iss["type"] == "error" for iss in issues)
    )

    return {
        "total": len(result_df),
        "instruments": instruments,
        "date_range": date_range,
        "columns": TEMPLATE_COLUMNS,
        "preview": preview_records,
        "data": result_df.fillna("").to_dict(orient="records"),
        "validation": {
            "total_errors": total_errors,
            "total_warnings": total_warnings,
            "rows_with_errors": rows_with_errors,
            "rows_with_warnings": rows_with_warnings,
            "issues": rows_detail,
        },
    }


# ── Excel export ───────────────────────────────────────────────

def _build_excel(data: list) -> io.BytesIO:
    df = pd.DataFrame(data, columns=TEMPLATE_COLUMNS)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    even_fill = PatternFill("solid", fgColor="FFFFFF")
    odd_fill = PatternFill("solid", fgColor="F2F2F2")

    for c, header in enumerate(TEMPLATE_COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, row_data in enumerate(df.itertuples(index=False), 2):
        fill = even_fill if r % 2 == 0 else odd_fill
        for c, value in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")

    for c, header in enumerate(TEMPLATE_COLUMNS, 1):
        col_vals = [str(header)] + [str(v) for v in df.iloc[:, c - 1]]
        width = min(max(len(s) for s in col_vals), 40) + 2
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


_SPLIT_CHUNK_SIZE = 5000


def _build_excel_split(data: list) -> io.BytesIO:
    chunks = [data[i:i + _SPLIT_CHUNK_SIZE] for i in range(0, len(data), _SPLIT_CHUNK_SIZE)]
    total_parts = len(chunks)
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for idx, chunk in enumerate(chunks, 1):
            excel_buf = _build_excel(chunk)
            zf.writestr(f"trade_report_part{idx}_of{total_parts}.xlsx", excel_buf.read())
    zip_buf.seek(0)
    return zip_buf


# ── Endpoints ──────────────────────────────────────────────────

@router.post("/process")
async def process_file(
    file: UploadFile = File(...),
    account: str = Form(...),
    exchange: str = Form(...),
    current_user: str = Depends(get_current_user),
):
    ext = "." + (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in _ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Допустимые форматы: .csv, .xlsx, .xls")

    filepath = None
    try:
        filepath = save_upload_file(file)
        result = await run_in_threadpool(
            _process_file, filepath, ext, account.strip(), exchange.strip()
        )
        result_id = uuid.uuid4().hex
        with _CACHE_LOCK:
            _CACHE[result_id] = {
                "owner": current_user,
                "data": result["data"],
                "created_at": datetime.now(),
            }
        return {
            "status": "success",
            "result_id": result_id,
            "total": result["total"],
            "instruments": result["instruments"],
            "date_range": result["date_range"],
            "columns": result["columns"],
            "preview": result["preview"],
            "validation": result["validation"],
        }
    except ValueError as e:
        raise HTTPException(400, str(e))
    except HTTPException:
        raise
    except Exception as e:
        log.error("convert-template error: %s", e, exc_info=True)
        raise HTTPException(500, "Ошибка обработки файла")
    finally:
        cleanup_files(filepath)


@router.get("/export/{result_id}")
async def export_excel(
    result_id: str,
    split: bool = Query(False),
    current_user: str = Depends(get_current_user),
):
    with _CACHE_LOCK:
        entry = _CACHE.get(result_id)
    if not entry:
        raise HTTPException(404, "Результат не найден или истёк срок хранения")
    if entry["owner"] != current_user:
        raise HTTPException(403, "Нет доступа")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    if split and len(entry["data"]) > _SPLIT_CHUNK_SIZE:
        buf = await run_in_threadpool(_build_excel_split, entry["data"])
        return StreamingResponse(
            buf,
            media_type="application/zip",
            headers={"Content-Disposition": f'attachment; filename="trade_report_{ts}.zip"'},
        )

    buf = await run_in_threadpool(_build_excel, entry["data"])
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="trade_report_{ts}.xlsx"'},
    )
